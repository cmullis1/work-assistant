"""Model test runs in the app: the test routes, and every results screen ignoring test reviews.

Review Focus 3 is `test_test_reviews_invisible_to_results`. Synthetic data only; no Zendesk or Anthropic.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

from factories import add_run
from test_queries import LEE, SAM, add_review, all_met, missing, version  # noqa: F401  (fixture)
from test_pipeline import make_contact
from work_assistant import pipeline
from work_assistant.web import app as webapp
from work_assistant.web.app import create_app

STATIC = Path(webapp.__file__).parent / "static"
BASE = "http://127.0.0.1:8765"
SONNET, FABLE = "claude-sonnet-5-5", "claude-fable-5-1"
WINDOW = (datetime(2026, 9, 20, 5, tzinfo=timezone.utc), datetime(2026, 10, 11, 5, tzinfo=timezone.utc))


def client(tmp_path, *, starter=None, zendesk_factory=None):
    def no_zendesk():
        raise AssertionError("a model test reads stored calls; it never asks Zendesk for anything")

    app = create_app(tmp_path / "work-assistant.db", run_starter=starter or (lambda spec: 4242),
                     zendesk_factory=zendesk_factory or no_zendesk, log_dir=tmp_path / "logs", startup_work=False)
    return TestClient(app, base_url=BASE)


def opposite(outcomes: dict[str, str]) -> dict[str, str]:
    return {cid: "met" if out != "met" else "missed" for cid, out in outcomes.items()}


def real_snapshots(conn) -> None:
    """Give every stored call a full snapshot, as the pipeline stores them."""
    for row in conn.execute("SELECT ticket_id, call_started_at FROM contacts WHERE purged_at IS NULL").fetchall():
        contact = make_contact(row["ticket_id"], call_started_at=row["call_started_at"])
        conn.execute("UPDATE contacts SET snapshot = ? WHERE ticket_id = ?",
                     (pipeline.contact_snapshot(contact), row["ticket_id"]))
    conn.commit()


def add_unreviewed_call(conn, ticket_id: int, call: str) -> None:
    conn.execute(
        "INSERT INTO contacts (ticket_id, channel, call_started_at, agent_email, agent_name, attribution_ok, snapshot, "
        "fetched_at) VALUES (?, 'voice', ?, ?, ?, 1, '{}', '2026-10-12T00:00:00Z')", (ticket_id, call, *SAM))
    conn.commit()


def add_test_review(conn, version, test_run: int, ticket_id: int, outcomes: dict[str, str], *, status="ok",
                    counted=True) -> int:
    """A test review of a call, paired with its shown review, with every override set too."""
    shown = conn.execute("SELECT review_id, rubric_version_id FROM reviews WHERE ticket_id = ? AND is_current = 1",
                         (ticket_id,)).fetchone()
    review_id = conn.execute(
        "INSERT INTO reviews (run_id, ticket_id, rubric_version_id, status, error, is_current, created_at, counted, "
        "model, effort, compared_review_id) VALUES (?, ?, ?, ?, ?, 0, '2026-10-13T00:00:00Z', ?, ?, 'high', ?)",
        (test_run, ticket_id, shown["rubric_version_id"] if shown else version.version_id, status,
         None if status == "ok" else "invalid output: test", int(counted), FABLE,
         shown["review_id"] if shown else None)).lastrowid
    conn.executemany(
        "INSERT INTO check_results (review_id, check_id, outcome, evidence, decided_by, override_outcome, "
        "needs_recheck) VALUES (?, ?, ?, 'test evidence', 'model', ?, 1)",
        [(review_id, cid, out, "unsure") for cid, out in outcomes.items()])
    conn.commit()
    return review_id


def sheet(content: bytes) -> list[tuple]:
    return list(openpyxl.load_workbook(io.BytesIO(content)).active.iter_rows(values_only=True))


def everything_shown(api) -> dict:
    """Every results endpoint's output."""
    shown = {}
    for path in ("/api/overview", f"/api/overview?agent={SAM[0]}", "/api/overview?week=2026-09-27",
                 "/api/agents", f"/api/agents/{SAM[0]}", f"/api/agents/{LEE[0]}?weeks=4",
                 "/api/tickets", "/api/tickets?verdict=not_counted", "/api/tickets?verdict=problem",
                 "/api/tickets?missed_check=sc06", "/api/tickets?q=frozen",
                 "/api/tickets/1", "/api/tickets/2", "/api/tickets/3", "/api/tickets/4", "/api/tickets/5",
                 "/api/agreement", "/api/agreement?weeks=52", "/api/rubric"):
        response = api.get(path)
        shown[path] = (response.status_code, response.json())
    for path in ("/api/export.xlsx", "/api/export.xlsx?verdict=not_counted"):
        shown[path] = sheet(api.get(path).content)
    return shown


# -- the brief's tests ----------------------------------------------------------------------------


def test_test_reviews_invisible_to_results(conn, version, tmp_path, monkeypatch):
    class Today(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(2026, 10, 12, 12, tzinfo=timezone.utc).astimezone(tz)

    monkeypatch.setattr(webapp, "datetime", Today)  # so the agreement period covers the calls
    shown_1 = add_review(conn, version, 1, missing(version, "sc06"), agent=SAM, call="2026-10-05T15:00:00Z",
                         overrides={"sc02": "missed"})
    add_review(conn, version, 2, all_met(version), agent=LEE, call="2026-10-06T15:00:00Z")
    add_review(conn, version, 3, missing(version, "sc15"), agent=SAM, call="2026-09-28T15:00:00Z", counted=False)
    add_review(conn, version, 4, {"pr-call-length": "met"}, agent=LEE, call="2026-09-21T15:00:00Z", status="error")
    add_unreviewed_call(conn, 5, "2026-10-07T15:00:00Z")
    real_snapshots(conn)
    api = client(tmp_path)
    before = everything_shown(api)
    assert before["/api/tickets/5"][0] == 404 and before["/api/agreement"][1]["overall"] is not None
    first_run = conn.execute("SELECT MIN(run_id) FROM runs").fetchone()[0]
    already = {t: pipeline._already_reviewed(conn, t, version.version_id) for t in range(1, 6)}
    later = {t: pipeline._reviewed_by_later_run(conn, t, first_run) for t in range(1, 6)}
    candidates = pipeline._rereview_candidates(conn, WINDOW, [], pipeline._QuietLog())

    # Opposite outcomes from a different model, overrides on every check, a Not counted flip, an error, and a
    # test of a call with no shown review at all (a review error under the current rubric version).
    test_run = add_run(conn, kind="test", status="complete", actual_usd=0.4)
    add_test_review(conn, version, test_run, 1, opposite(missing(version, "sc06")), counted=False)
    add_test_review(conn, version, test_run, 2, opposite(all_met(version)))
    add_test_review(conn, version, test_run, 3, opposite(missing(version, "sc15")))
    add_test_review(conn, version, test_run, 4, opposite(all_met(version)))
    add_test_review(conn, version, test_run, 5, {"pr-call-length": "missed"}, status="error")
    paused_test = add_run(conn, kind="test", status="awaiting_confirmation")
    conn.execute("INSERT INTO run_contacts (run_id, ticket_id) VALUES (?, 2)", (paused_test,))
    conn.commit()

    assert everything_shown(api) == before
    assert {t: pipeline._already_reviewed(conn, t, version.version_id) for t in range(1, 6)} == already
    assert {t: pipeline._reviewed_by_later_run(conn, t, first_run) for t in range(1, 6)} == later
    assert pipeline._rereview_candidates(conn, WINDOW, [], pipeline._QuietLog()) == candidates
    assert api.get("/api/tickets/1").json()["review"]["review_id"] == shown_1


def test_api_test_estimate_and_start(conn, version, tmp_path):
    add_review(conn, version, 1, all_met(version), call="2026-10-05T15:00:00Z")
    add_review(conn, version, 2, all_met(version), call="2026-10-06T15:00:00Z", status="error")
    conn.execute(
        "INSERT INTO contacts (ticket_id, channel, call_started_at, agent_email, agent_name, attribution_ok, snapshot, "
        "fetched_at, purged_at) VALUES (3, 'voice', '2026-10-07T15:00:00Z', 'a@x.com', 'A', 1, '{}', 'x', 'y')")
    conn.commit()
    real_snapshots(conn)
    started = []
    api = client(tmp_path, starter=lambda spec: started.append(spec) or 4242)
    body = {"from": "2026-10-04", "to": "2026-10-10", "model": SONNET, "effort": "high"}

    estimate = api.post("/api/tests/estimate", json=body)
    assert estimate.status_code == 200
    result = estimate.json()
    assert result["estimate_usd"] > 0
    assert {k: v for k, v in result.items() if k != "estimate_usd"} == {
        "contacts": 1, "excluded_purged": 1, "excluded_errored": 1, "excluded_unreviewed": 0}
    assert started == []
    assert conn.execute("SELECT COUNT(*) FROM runs WHERE kind = 'test'").fetchone()[0] == 0
    assert api.post("/api/tests/estimate", json={**body, "limit": 1}).json()["contacts"] == 1

    response = api.post("/api/tests", json={**body, "limit": 10})
    assert response.status_code == 202 and response.json() == {"started": True, "pid": 4242}
    assert api.post("/api/tests", json=body).status_code == 202
    assert started == [
        {"command": "test-model", "from": "2026-10-04", "to": "2026-10-10", "model": SONNET, "effort": "high",
         "limit": 10},
        {"command": "test-model", "from": "2026-10-04", "to": "2026-10-10", "model": SONNET, "effort": "high",
         "limit": None},
    ]

    bad = [
        ({**body, "model": "claude-nope"}, "Claude Sonnet 5.5"),
        ({k: v for k, v in body.items() if k != "model"}, "model"),
        ({**body, "effort": "extreme"}, "low, medium, high, xhigh, max"),
        ({**body, "from": "2026-10-11"}, "on or before"),
        ({**body, "to": "10/10/2026"}, "YYYY-MM-DD"),
        ({**body, "limit": 0}, "at least 1"),
        ({**body, "limit": "5"}, "at least 1"),
        ({**body, "limit": True}, "at least 1"),
        ({**body, "limit": 2.5}, "at least 1"),
        ([1, 2], "range"),
    ]
    for payload, message in bad:
        for path in ("/api/tests/estimate", "/api/tests"):
            response = api.post(path, json=payload)
            assert response.status_code == 400, (path, payload)
            assert message in response.json()["detail"], (path, payload)
    assert len(started) == 2

    for path in ("/api/tests/estimate", "/api/tests"):
        assert api.post(path, json=body, headers={"origin": "https://evil.example"}).status_code == 403
        assert api.post(path, json=body, headers={"host": "evil.example:8765"}).status_code == 403
    assert len(started) == 2


# -- more -----------------------------------------------------------------------------------------


def test_cli_args_for_test_model(tmp_path):
    db_file = tmp_path / "x.db"
    spec = {"command": "test-model", "from": "2026-10-04", "to": "2026-10-10", "model": SONNET, "effort": "max",
            "limit": 10}
    args = webapp.cli_args(spec, db_file)
    # No --yes: the run's own estimate goes through the budget guard, which pauses it for the banner.
    assert args[1:] == ["-m", "work_assistant.cli", "--db", str(db_file), "test-model", "--from", "2026-10-04",
                        "--to", "2026-10-10", "--model", SONNET, "--effort", "max", "--limit", "10"]
    assert "--limit" not in webapp.cli_args({**spec, "limit": None}, db_file)


def test_runs_list_shows_test_runs(conn, version, tmp_path):
    add_run(conn, kind="test", status="awaiting_confirmation", estimate_usd=55.0)
    api = client(tmp_path)
    assert [r["kind"] for r in api.get("/api/runs").json()["runs"]] == ["test"]
    assert [(r["run_id"], r["kind"]) for r in api.get("/api/status").json()["awaiting_confirmation"]] == [(1, "test")]


def test_kind_label_for_model_tests():
    js = (STATIC / "app.js").read_text(encoding="utf-8")
    start = js.index("var KIND_LABELS")
    labels = js[start:js.index(";", start)]
    assert 'test: "Model test"' in labels
    assert 'scheduled: "Daily"' in labels


def test_scheduled_run_label_follows_its_window():
    """Run history calls a scheduled run Daily or Weekly by its window, so the weekly schedule's runs (and the
    demo's) still read Weekly after the switch to daily runs, even across a DST change."""
    import json as json_

    from test_web_overview import js_function, run_node

    js = (STATIC / "app.js").read_text(encoding="utf-8")
    labels = js[js.index("var KIND_LABELS"):js.index(";", js.index("var KIND_LABELS")) + 1]
    runs = [
        {"kind": "scheduled", "window_start": "2026-10-18T05:00:00Z", "window_end": "2026-10-19T05:00:00Z"},
        {"kind": "scheduled", "window_start": "2026-11-01T05:00:00Z", "window_end": "2026-11-02T06:00:00Z"},  # 25 h
        {"kind": "scheduled", "window_start": "2026-10-11T05:00:00Z", "window_end": "2026-10-18T05:00:00Z"},
        {"kind": "scheduled", "window_start": "2026-03-08T06:00:00Z", "window_end": "2026-03-15T05:00:00Z"},  # 7 d - 1 h
        {"kind": "scheduled", "window_start": "2026-10-12T05:00:00Z", "window_end": "2026-10-18T05:00:00Z"},  # 6 days
        {"kind": "manual", "window_start": "2026-10-11T05:00:00Z", "window_end": "2026-10-18T05:00:00Z"},
        {"kind": "test", "window_start": "2026-10-18T05:00:00Z", "window_end": "2026-10-19T05:00:00Z"},
    ]
    script = (labels + "\n" + js_function(js, "kindLabel") + "\n" +
              f"process.stdout.write(JSON.stringify({json_.dumps(runs)}.map(kindLabel)));")
    assert run_node(script) == ["Daily", "Daily", "Weekly", "Weekly", "Daily", "Run now", "Model test"]
    assert "kindLabel(run)" in js_function(js, "appendRun")


def test_status_last_run_ignores_test_runs(conn, version, tmp_path):
    api = client(tmp_path)
    add_run(conn, kind="test", status="failed", error="ZendeskError")
    assert api.get("/api/status").json()["last_run"] is None   # only a test so far: nothing to report

    weekly = add_run(conn, kind="scheduled", status="failed", error="ZendeskError")
    add_run(conn, kind="test", status="complete")
    add_run(conn, kind="test", status="failed", error=pipeline.INTERRUPTED)
    last = api.get("/api/status").json()["last_run"]
    assert (last["run_id"], last["status"], last["alert"]) == (weekly, "failed", True)  # a test can't hide it


def test_test_failure_labels_point_to_test_a_model(conn, version, tmp_path):
    for error in (pipeline.INTERRUPTED, "ZendeskError", "ConnectTimeout"):
        label = webapp.failure_label(error, "test")
        assert "Run now" not in label, error
        assert "Settings → Test a model" in label and "work-assistant test-model" in label, error
    assert webapp.failure_label(pipeline.INTERRUPTED, "manual") == webapp.RUN_INTERRUPTED
    assert webapp.failure_label("APITimeoutError", "test") == webapp.ANTHROPIC_UNREACHABLE

    add_run(conn, kind="test", status="failed", error=pipeline.INTERRUPTED)
    add_run(conn, kind="manual", status="failed", error=pipeline.INTERRUPTED)
    runs = client(tmp_path).get("/api/runs").json()["runs"]
    assert [r["error_label"] for r in runs] == [webapp.RUN_INTERRUPTED, webapp.TEST_INTERRUPTED]
