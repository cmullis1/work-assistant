"""The Excel export (web/export.py). Synthetic data only."""
from __future__ import annotations

import io

import openpyxl

from work_assistant.web import export

HEADER = ["Ticket", "Date", "Tech", "Priority", "Score", "Percent", "Verdict", "Missed", "Zendesk link"]


def row(ticket_id: int, **changes) -> dict:
    base = {
        "ticket_id": ticket_id, "date": "2026-10-05", "agent_email": "sam@example.com", "agent_name": "Sam Agent",
        "priority": "high", "earned": 90, "total": 100, "percent": 90.0, "verdict": "problem", "counted": True,
        "missed": ["#6", "10-minute rule"], "missed_ids": ["sc06", "pr-call-length"], "needs_recheck": False,
        "review_error": False, "error": None, "not_counted_reason": None,
    }
    return {**base, **changes}


def sheet_rows(rows, subdomain="acme") -> list[tuple]:
    buffer = io.BytesIO()
    export.write_xlsx(rows, buffer, subdomain=subdomain)
    return list(openpyxl.load_workbook(io.BytesIO(buffer.getvalue())).active.iter_rows(values_only=True))


def test_export_columns():
    rows = [row(1), row(2, verdict="good", missed=[], missed_ids=[], earned=100, percent=100.0), row(3)]
    sheet = sheet_rows(rows)
    assert list(sheet[0]) == HEADER
    assert len(sheet) == 1 + len(rows)
    first = sheet[1]
    assert first[0] == 1 and first[1].date().isoformat() == "2026-10-05"
    assert first[2:8] == ("Sam Agent", "High", 90, 90.0, "Problem", "#6, 10-minute rule")
    assert first[8] == "https://acme.zendesk.com/agent/tickets/1"
    assert sheet[2][6] == "Good" and sheet[2][7] is None


def test_export_review_error_and_not_counted_have_no_score():
    sheet = sheet_rows([row(1, review_error=True, verdict="review_failed"),
                        row(2, counted=False, verdict="not_counted", earned=None, total=None, percent=None,
                            missed=[], missed_ids=[])])
    assert sheet[1][4:7] == (None, None, "Review failed")
    assert sheet[2][4:7] == (None, None, "Not counted")


def test_export_link_is_a_hyperlink_and_blank_without_subdomain():
    buffer = io.BytesIO()
    export.write_xlsx([row(5)], buffer, subdomain="acme")
    cell = openpyxl.load_workbook(io.BytesIO(buffer.getvalue())).active["I2"]
    assert cell.hyperlink is not None and cell.hyperlink.target == "https://acme.zendesk.com/agent/tickets/5"
    assert sheet_rows([row(5)], subdomain=None)[1][8] is None


def test_export_never_writes_formulas():
    """Agent names and priorities come from Zendesk; a leading "=" must stay text, never a formula."""
    buffer = io.BytesIO()
    export.write_xlsx([row(1, agent_name='=HYPERLINK("http://evil.example","x")', priority="=1+1")], buffer,
                      subdomain="acme")
    sheet = openpyxl.load_workbook(io.BytesIO(buffer.getvalue())).active
    assert sheet["C2"].data_type == "s" and sheet["C2"].value.startswith("=HYPERLINK")
    assert sheet["D2"].data_type == "s"


def test_export_to_path(tmp_path):
    path = tmp_path / "out.xlsx"
    export.write_xlsx([row(1)], path, subdomain=None)
    assert [c.value for c in openpyxl.load_workbook(path).active[1]] == HEADER
