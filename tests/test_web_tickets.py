"""The Tickets screen endpoints, overrides ("Disagree"), the Excel export route, and page assets.
Synthetic data only."""
from __future__ import annotations

import io
import json
import re
import shutil
import subprocess
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from test_queries import LEE, SAM, add_review, all_met, missing, version  # noqa: F401  (fixture)
from test_web_overview import js_function, js_var, node_prelude, run_node
from work_assistant.models import Comment, FirstContact, MergeInfo
from work_assistant.pipeline import contact_snapshot
from work_assistant.retention import purge_expired
from work_assistant.settings import set_setting
from work_assistant.web import app as webapp
from work_assistant.web.app import create_app

from factories import add_run

STATIC = Path(webapp.__file__).parent / "static"
BASE = "http://127.0.0.1:8765"
SAME_ORIGIN = {"origin": BASE}
CALL = "2026-10-05T15:00:00Z"          # a Monday, 10:00 in America/Chicago
TIMELINE_KEYS = {"call_start", "call_end", "triage_macro_at", "call_limit_at", "summary_email_at",
                 "email_deadline_at", "callback_at", "merged_into"}
CHECK_KEYS = {"check_id", "kind", "number", "summary", "points", "outcome", "evidence", "decided_by",
              "override_outcome", "override_note", "needs_recheck"}


def client(tmp_path):
    app = create_app(tmp_path / "work-assistant.db", run_starter=lambda spec: 4242,
                     zendesk_factory=lambda: None, log_dir=tmp_path / "logs", startup_work=False)
    return TestClient(app, base_url=BASE)


def comment(n: int, at: str, kind: str, text: str, role: str = "agent", public: bool = False) -> Comment:
    return Comment(comment_id=n, created_at=at, author_id=7 if role == "agent" else 0,
                   author_name=SAM[1] if role == "agent" else "System", author_role=role, public=public,
                   kind=kind, text=text, image_urls=())


def make_call(ticket_id: int, *, priority: str = "normal", start: str = CALL, duration: int = 720,
              email_at: str | None = "2026-10-05T18:00:00Z", callback_minutes: float | None = None,
              merged: MergeInfo | None = None, agent=SAM) -> FirstContact:
    """A 12-minute inbound call: transcript, triage macro at 15:03, summary email at 18:00."""
    comments = [
        comment(1, "2026-10-05T15:12:00Z", "transcript", "[00:03] Agent: Thanks for calling.\n[00:20] Caller: "
                "<b>help</b> the receipt printer is offline", role="system"),
        comment(2, "2026-10-05T15:03:00Z", "triage_macro", "Location: Demo Store 4\nIssue: receipt printer"),
    ]
    if email_at is not None:
        comments.append(comment(3, email_at, "summary_email", "Hello, here is a summary.", public=True))
    return FirstContact(
        ticket_id=ticket_id, channel="voice", subject="Receipt printer offline", priority=priority, tags=("pos",),
        fields={"Location ID": "DEMO-104"}, call_started_at=start, call_duration_seconds=duration,
        agent_user_id=7, agent_email=agent[0], agent_name=agent[1], attribution_ok=True, summary_email_at=email_at,
        comments=tuple(comments), callback_minutes=callback_minutes, merged_into=merged,
    )


def add_call(conn, version, contact: FirstContact, outcomes: dict[str, str], *, status: str = "ok",
             error: str | None = None, counted: bool = True, recheck=()) -> int:
    """A stored call with a full snapshot and its current review (timed checks decided by code)."""
    run_id = conn.execute("SELECT MIN(run_id) FROM runs").fetchone()[0] or add_run(conn)
    conn.execute(
        "INSERT INTO contacts (ticket_id, channel, call_started_at, agent_email, agent_name, priority, "
        "attribution_ok, snapshot, fetched_at) VALUES (?, 'voice', ?, ?, ?, ?, 1, ?, '2026-10-12T00:00:00Z')",
        (contact.ticket_id, contact.call_started_at, contact.agent_email, contact.agent_name, contact.priority,
         contact_snapshot(contact)))
    review_id = conn.execute(
        "INSERT INTO reviews (run_id, ticket_id, rubric_version_id, status, error, resolved_on_call, counted, "
        "not_counted_reason, is_current, created_at) VALUES (?, ?, ?, ?, ?, 'yes', ?, ?, 1, '2026-10-12T00:00:00Z')",
        (run_id, contact.ticket_id, version.version_id, status, error, int(counted),
         None if counted else "Wrong number")).lastrowid
    timed = {c.check_id for c in version.checks if c.timed}
    conn.executemany(
        "INSERT INTO check_results (review_id, check_id, outcome, evidence, decided_by, needs_recheck) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        [(review_id, cid, out, f"evidence for {cid}", "code" if cid in timed else "model", int(cid in recheck))
         for cid, out in outcomes.items()])
    conn.commit()
    return review_id


@pytest.fixture
def call(conn, version):
    set_setting(conn, "zendesk_subdomain", "acme")
    review_id = add_call(conn, version, make_call(101), missing(version, "sc06"))
    return review_id


def rows(api, query: str = "") -> list[dict]:
    response = api.get("/api/tickets" + query)
    assert response.status_code == 200, response.text
    return response.json()["rows"]


def check(detail: dict, check_id: str) -> dict:
    return {c["check_id"]: c for c in detail["checks"]}[check_id]


# -- the brief's tests ----------------------------------------------------------------------------------


def test_ticket_detail_shape(call, version, tmp_path):
    response = client(tmp_path).get("/api/tickets/101")
    assert response.status_code == 200
    body = response.json()
    assert set(body) >= {"contact", "review", "timeline", "checks", "comments", "zendesk_url"}
    assert set(body["timeline"]) >= TIMELINE_KEYS
    # The plan said 21 checks; the seeded rubric has since gained two process rules (dropped-call
    # callback, merge), so a non-urgent call lists every enabled check of its version: 23.
    enabled = [c for c in version.checks if c.enabled]
    assert len(enabled) == 23
    assert len(body["checks"]) == len(enabled)
    for item in body["checks"]:
        assert set(item) >= CHECK_KEYS
    review = body["review"]
    assert review["review_id"] == call and review["version_id"] == version.version_id
    assert (review["earned"], review["total"], review["percent"], review["verdict"]) == (95, 100, 95.0, "problem")
    assert review["resolved_on_call"] == "yes" and review["error"] is None
    assert body["zendesk_url"] == "https://acme.zendesk.com/agent/tickets/101"
    assert [c["kind"] for c in body["comments"]] == ["triage_macro", "transcript", "summary_email"]  # by time


def test_override_roundtrip(call, tmp_path):
    api = client(tmp_path)
    assert rows(api)[0]["verdict"] == "problem"
    url = f"/api/reviews/{call}/checks/sc06/override"
    response = api.put(url, json={"outcome": "met", "note": "She did summarize it at 4:10."}, headers=SAME_ORIGIN)
    assert response.status_code == 200, response.text
    sc06 = check(response.json(), "sc06")
    assert (sc06["outcome"], sc06["override_outcome"], sc06["override_note"]) == \
        ("missed", "met", "She did summarize it at 4:10.")
    [row] = rows(api)
    assert (row["verdict"], row["percent"], row["missed"]) == ("good", 100.0, [])

    assert api.delete(url, headers=SAME_ORIGIN).status_code == 200
    [row] = rows(api)
    assert (row["verdict"], row["percent"], row["missed"]) == ("problem", 95.0, ["#6"])


def test_override_invalid_outcome_400(call, tmp_path):
    api = client(tmp_path)
    url = f"/api/reviews/{call}/checks/sc06/override"
    for body in ({"outcome": "great"}, {"outcome": ""}, {"note": "no outcome"}, {"outcome": "MET"}, ["met"]):
        assert api.put(url, json=body, headers=SAME_ORIGIN).status_code == 400, body
    assert api.put(url, json={"outcome": "met", "note": 5}, headers=SAME_ORIGIN).status_code == 400
    assert check(api.get("/api/tickets/101").json(), "sc06")["override_outcome"] is None


# -- overrides ------------------------------------------------------------------------------------------


def test_override_clears_recheck_and_undo_writes_null(conn, version, tmp_path):
    review_id = add_call(conn, version, make_call(102), missing(version, "sc15"), recheck=("sc15",))
    api = client(tmp_path)
    assert check(api.get("/api/tickets/102").json(), "sc15")["needs_recheck"] is True
    url = f"/api/reviews/{review_id}/checks/sc15/override"
    api.put(url, json={"outcome": "na", "note": "  "}, headers=SAME_ORIGIN)
    row = conn.execute("SELECT * FROM check_results WHERE review_id = ? AND check_id = 'sc15'", (review_id,)).fetchone()
    assert (row["override_outcome"], row["override_note"], row["needs_recheck"]) == ("na", None, 0)
    assert datetime.fromisoformat(row["override_at"].replace("Z", "+00:00")).tzinfo is not None

    detail = api.delete(url, headers=SAME_ORIGIN).json()
    assert check(detail, "sc15")["override_outcome"] is None
    row = conn.execute("SELECT * FROM check_results WHERE review_id = ? AND check_id = 'sc15'", (review_id,)).fetchone()
    assert (row["override_outcome"], row["override_note"], row["override_at"]) == (None, None, None)


def test_override_reflected_in_overview(call, tmp_path):
    api = client(tmp_path)
    assert api.get("/api/overview").json()["tiles"]["problems"]["count"] == 1
    api.put(f"/api/reviews/{call}/checks/sc06/override", json={"outcome": "met"}, headers=SAME_ORIGIN)
    assert api.get("/api/overview").json()["tiles"]["problems"]["count"] == 0


def test_override_unknown_review_or_check_404(call, tmp_path):
    api = client(tmp_path)
    assert api.put("/api/reviews/999/checks/sc06/override", json={"outcome": "met"},
                   headers=SAME_ORIGIN).status_code == 404
    assert api.put(f"/api/reviews/{call}/checks/nope/override", json={"outcome": "met"},
                   headers=SAME_ORIGIN).status_code == 404
    assert api.delete(f"/api/reviews/{call}/checks/nope/override", headers=SAME_ORIGIN).status_code == 404


def test_override_on_replaced_review_409(call, conn, tmp_path):
    conn.execute("UPDATE reviews SET is_current = 0 WHERE review_id = ?", (call,))
    conn.commit()
    response = client(tmp_path).put(f"/api/reviews/{call}/checks/sc06/override", json={"outcome": "met"},
                                    headers=SAME_ORIGIN)
    assert response.status_code == 409


@pytest.mark.parametrize("origin", ["http://evil.example", "http://127.0.0.1:9999"])
def test_override_from_a_foreign_origin_rejected(call, tmp_path, origin):
    api = client(tmp_path)
    url = f"/api/reviews/{call}/checks/sc06/override"
    assert api.put(url, json={"outcome": "met"}, headers={"origin": origin}).status_code == 403
    assert api.delete(url, headers={"origin": origin}).status_code == 403


def test_override_note_too_long_400(call, tmp_path):
    response = client(tmp_path).put(f"/api/reviews/{call}/checks/sc06/override",
                                    json={"outcome": "met", "note": "x" * 10001}, headers=SAME_ORIGIN)
    assert response.status_code == 400


# -- detail ---------------------------------------------------------------------------------------------


def test_timeline_values(conn, version, tmp_path):
    merged = MergeInfo(ticket_id=900, subject="Same site, earlier ticket", location="DEMO-104",
                       created_at="2026-10-04T12:00:00Z",
                       comments=(comment(9, "2026-10-05T17:00:00Z", "summary_email", "Merged reply", public=True),))
    add_call(conn, version, make_call(103, callback_minutes=5, merged=merged), all_met(version))
    set_setting(conn, "zendesk_subdomain", "acme")
    body = client(tmp_path).get("/api/tickets/103").json()
    t = body["timeline"]
    assert t["call_start"] == "2026-10-05T15:00:00Z"
    assert t["call_end"] == "2026-10-05T15:12:00Z"
    assert t["triage_macro_at"] == "2026-10-05T15:03:00Z"
    assert t["call_limit_at"] == "2026-10-05T15:10:00Z"
    assert t["summary_email_at"] == "2026-10-05T18:00:00Z"
    assert t["email_deadline_at"] == "2026-10-05T15:42:00Z"   # the seed's 30 minutes after the call ends
    assert t["email_deadline_minutes"] == 30
    assert "email_deadline_hours" not in t
    assert t["callback_at"] == "2026-10-05T15:17:00Z"
    assert t["merged_into"] == {"ticket_id": 900, "subject": "Same site, earlier ticket",
                                "zendesk_url": "https://acme.zendesk.com/agent/tickets/900"}
    assert [c["text"] for c in body["merged_comments"]] == ["Merged reply"]


def test_urgent_call_has_no_call_limit(conn, version, tmp_path):
    add_call(conn, version, make_call(104, priority="urgent", email_at=None), all_met(version))
    t = client(tmp_path).get("/api/tickets/104").json()["timeline"]
    assert t["call_limit_at"] is None and t["exempt"] is True
    assert t["summary_email_at"] is None and t["email_deadline_at"] == "2026-10-05T15:42:00Z"
    assert t["callback_at"] is None and t["merged_into"] is None


def test_check_fields(call, tmp_path):
    body = client(tmp_path).get("/api/tickets/101").json()
    sc06 = check(body, "sc06")
    assert (sc06["kind"], sc06["number"], sc06["points"], sc06["outcome"], sc06["earned_points"]) == \
        ("scorecard", 6, 5, "missed", 0)
    assert sc06["decided_by"] == "model" and sc06["evidence"] == "evidence for sc06"
    length = check(body, "pr-call-length")
    assert (length["kind"], length["number"], length["decided_by"], length["earned_points"]) == \
        ("process", None, "code", None)
    assert [c["check_id"] for c in body["checks"]][:3] == ["sc01", "sc02", "sc03"]


def test_comment_text_is_returned_as_plain_text(call, tmp_path):
    body = client(tmp_path).get("/api/tickets/101").json()
    transcript = [c for c in body["comments"] if c["kind"] == "transcript"][0]
    assert "<b>help</b>" in transcript["text"]       # untouched; the page renders it with textContent
    assert set(transcript) >= {"created_at", "author_name", "author_role", "kind", "public", "text"}


def test_review_error_has_no_score(conn, version, tmp_path):
    add_call(conn, version, make_call(105), {"pr-call-length": "missed", "pr-email-sent": "met"}, status="error",
             error="output cut off at max_tokens")
    api = client(tmp_path)
    body = api.get("/api/tickets/105").json()
    review = body["review"]
    assert review["review_error"] is True and review["error"] == "output cut off at max_tokens"
    assert review["earned"] is None and review["percent"] is None
    assert review["verdict"] == "review_failed"   # its own verdict, even with a measured miss (T-045)
    [row] = rows(api)
    assert row["review_error"] is True and row["error"] == "output cut off at max_tokens"
    assert row["earned"] is None and row["percent"] is None   # never a 0 score
    assert row["verdict"] == "review_failed"


def test_review_failed_filter(conn, version, tmp_path):
    """T-045: the Verdict filter lists the calls whose review did not finish, and only those; Problem no
    longer hides them, and the scores of the other calls don't change."""
    add_call(conn, version, make_call(140), missing(version, "sc06"))
    add_call(conn, version, make_call(141, start="2026-10-06T15:00:00Z"), {"pr-call-length": "met"},
             status="error", error="request errored: overloaded_error")
    add_call(conn, version, make_call(142, start="2026-10-07T15:00:00Z"), all_met(version))
    api = client(tmp_path)
    assert [r["ticket_id"] for r in rows(api, "?verdict=review_failed")] == [141]
    assert [r["ticket_id"] for r in rows(api, "?verdict=problem")] == [140]
    assert [r["ticket_id"] for r in rows(api, "?verdict=needs_review")] == []
    assert [r["ticket_id"] for r in rows(api)] == [142, 141, 140]   # an ordinary counted call in the full list
    overview = api.get("/api/overview").json()
    assert overview["tiles"]["avg_score"]["avg_percent"] == 97.5    # (95 + 100) / 2: the failed call has no score
    assert overview["tiles"]["calls"]["count"] == 3
    response = api.get("/api/tickets?verdict=awful")
    assert response.status_code == 400 and "review_failed" in response.json()["detail"]


def test_review_error_unjudged_checks_have_no_earned_points(conn, version, tmp_path):
    add_call(conn, version, make_call(108), {"pr-call-length": "met", "sc02": "missed"}, status="error",
             error="output cut off at max_tokens")
    body = client(tmp_path).get("/api/tickets/108").json()
    sc01 = check(body, "sc01")
    assert (sc01["outcome"], sc01["effective"], sc01["earned_points"]) == (None, None, None)  # "-", not 0 / N
    assert check(body, "sc02")["earned_points"] == 0   # a stored miss still shows 0 / N


def _shown_errors(api, ticket_id):
    """The error text the list row and the call detail show for one call."""
    [row] = [r for r in rows(api) if r["ticket_id"] == ticket_id]
    return row["error"], api.get(f"/api/tickets/{ticket_id}").json()["review"]["error"]


def test_refusal_message_plain(conn, version, tmp_path):
    add_call(conn, version, make_call(120), {"pr-call-length": "met"}, status="error", error="refused: cyber")
    message = "The model declined to review this call."
    assert _shown_errors(client(tmp_path), 120) == (message, message)


def test_fable_retention_error_plain(conn, version, tmp_path):
    error = "request errored: invalid_request_error"
    fable = add_call(conn, version, make_call(121), {"pr-call-length": "met"}, status="error", error=error)
    opus = add_call(conn, version, make_call(122, start="2026-10-06T15:00:00Z"), {"pr-call-length": "met"},
                    status="error", error=error)
    conn.execute("UPDATE reviews SET model = 'claude-fable-5-1' WHERE review_id = ?", (fable,))
    conn.execute("UPDATE reviews SET model = 'claude-opus-5-5' WHERE review_id = ?", (opus,))
    conn.commit()
    api = client(tmp_path)
    message = "Claude Fable 5.1 needs 30-day data retention, which your Anthropic organization doesn't allow."
    assert _shown_errors(api, 121) == (message, message)
    assert _shown_errors(api, 122) == (error, error)   # other models keep the stored text


def test_not_counted_hidden_by_default_and_listed_with_reason(conn, version, tmp_path):
    add_call(conn, version, make_call(106), all_met(version))
    add_call(conn, version, make_call(107, start="2026-10-06T15:00:00Z"), all_met(version), counted=False)
    api = client(tmp_path)
    assert [r["ticket_id"] for r in rows(api)] == [106]
    [row] = rows(api, "?verdict=not_counted")
    assert (row["ticket_id"], row["counted"], row["not_counted_reason"]) == (107, False, "Wrong number")
    detail = api.get("/api/tickets/107").json()
    assert detail["review"]["counted"] is False and detail["review"]["not_counted_reason"] == "Wrong number"
    assert detail["review"]["earned"] is None and detail["review"]["verdict"] == "not_counted"


def test_purged_call_keeps_scores_but_not_content(call, conn, tmp_path):
    set_setting(conn, "retention_days", 30)
    assert purge_expired(conn, datetime(2027, 1, 1, tzinfo=timezone.utc), 30, tmp_path / "logs") == 1
    body = client(tmp_path).get("/api/tickets/101").json()
    assert body["contact"]["purged"] is True and body["retention_days"] == 30
    assert body["comments"] == [] and body["merged_comments"] == []
    assert all(c["evidence"] is None for c in body["checks"])
    assert check(body, "sc06")["outcome"] == "missed"
    assert (body["review"]["earned"], body["review"]["verdict"]) == (95, "problem")
    assert body["timeline"]["call_start"] == CALL


def test_ticket_detail_404(conn, version, tmp_path):
    assert client(tmp_path).get("/api/tickets/4040").status_code == 404


def test_zendesk_url_none_without_subdomain(conn, version, tmp_path):
    add_call(conn, version, make_call(108), all_met(version))
    assert client(tmp_path).get("/api/tickets/108").json()["zendesk_url"] is None


# -- the list -------------------------------------------------------------------------------------------


def test_list_filters(conn, version, tmp_path):
    add_review(conn, version, 1, missing(version, "sc06"), agent=SAM, priority="high", subject="Kiosk frozen")
    add_review(conn, version, 2, all_met(version), agent=LEE, call="2026-10-06T15:00:00Z")
    add_review(conn, version, 3, all_met(version), agent=LEE, call="2026-09-29T15:00:00Z")
    api = client(tmp_path)
    assert [r["ticket_id"] for r in rows(api)] == [2, 1, 3]
    assert [r["ticket_id"] for r in rows(api, "?agent=LEE@example.com")] == [2, 3]
    assert [r["ticket_id"] for r in rows(api, "?missed_check=sc06")] == [1]
    assert [r["ticket_id"] for r in rows(api, "?verdict=good")] == [2, 3]
    assert [r["ticket_id"] for r in rows(api, "?start=2026-10-04&end=2026-10-10")] == [2, 1]
    assert [r["ticket_id"] for r in rows(api, "?q=kiosk")] == [1]
    assert [r["ticket_id"] for r in rows(api, "?priority=high")] == [1]
    assert [r["ticket_id"] for r in rows(api, "?verdict=&agent=&q=")] == [2, 1, 3]


def test_list_filter_options(call, conn, version, tmp_path):
    body = client(tmp_path).get("/api/tickets").json()
    assert {"email": SAM[0], "name": SAM[1]} in body["agent_options"]
    labels = {o["check_id"]: o["label"] for o in body["check_options"]}
    assert labels["sc06"] == "#6 Triage: Verify understanding"
    assert body["priorities"] == ["normal"]
    assert body["total"] == 1


@pytest.mark.parametrize("query", ["?verdict=awful", "?start=2026-13-01", "?end=yesterday"])
def test_list_rejects_bad_filters(call, tmp_path, query):
    assert client(tmp_path).get("/api/tickets" + query).status_code == 400


# -- export ---------------------------------------------------------------------------------------------


def read_sheet(content: bytes) -> list[tuple]:
    sheet = openpyxl.load_workbook(io.BytesIO(content)).active
    return list(sheet.iter_rows(values_only=True))


def test_export_route(conn, version, tmp_path):
    set_setting(conn, "zendesk_subdomain", "acme")
    add_call(conn, version, make_call(110), missing(version, "sc06"))
    add_call(conn, version, make_call(111, start="2026-10-06T15:00:00Z"), all_met(version), counted=False)
    add_call(conn, version, make_call(112, start="2026-10-07T15:00:00Z", agent=LEE),
             {"pr-call-length": "met"}, status="error", error="request errored: overloaded_error")
    api = client(tmp_path)
    before = sorted(p.name for p in tmp_path.rglob("*"))
    response = api.get("/api/export.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert re.match(r'attachment; filename="work-assistant-tickets-[0-9-]+\.xlsx"', response.headers["content-disposition"])
    assert sorted(p.name for p in tmp_path.rglob("*")) == before     # nothing written on the server side
    sheet = read_sheet(response.content)
    assert [r[0] for r in sheet[1:]] == [112, 110]                    # Not counted left out by default
    error_row, problem_row = sheet[1], sheet[2]
    assert (error_row[4], error_row[5], error_row[6]) == (None, None, "Review failed")
    assert problem_row[4:8] == (95, 95.0, "Problem", "#6")
    assert problem_row[8] == "https://acme.zendesk.com/agent/tickets/110"

    only = read_sheet(api.get("/api/export.xlsx?verdict=not_counted").content)
    assert [(r[0], r[6]) for r in only[1:]] == [(111, "Not counted")]
    assert [r[0] for r in read_sheet(api.get(f"/api/export.xlsx?agent={LEE[0]}").content)[1:]] == [112]
    assert api.get("/api/export.xlsx?verdict=awful").status_code == 400


# -- page assets ----------------------------------------------------------------------------------------


def test_static_assets_are_revalidated(tmp_path):
    api = client(tmp_path)
    for name in ("app.js", "app.css", "charts.js"):
        response = api.get(f"/static/{name}")
        assert response.status_code == 200
        assert response.headers["cache-control"] == "no-cache", name


def test_tickets_page_wired_up():
    js = (STATIC / "app.js").read_text(encoding="utf-8")
    assert "renderTickets" in js and "/api/tickets" in js and "/override" in js and "/api/export.xlsx" in js
    for text in ("Disagree", "Undo my override", "Review failed", "Not counted", "Details removed after",
                 "Open in Zendesk", "Recheck"):
        assert text in js, text
    for param in ("ticket", "verdict", "missed_check", "start", "end", "agent", "q"):
        assert '"' + param + '"' in js, param
    assert "innerHTML" not in js


def test_tickets_page_styles():
    css = (STATIC / "app.css").read_text(encoding="utf-8")
    for selector in (".timeline-strip", ".timeline-chip", ".dethead", ".comment-body", ".atable tr.sel td",
                     ".override-form", ".badge-recheck"):
        assert selector in css, selector
    body = re.search(r"\.comment-body\s*\{[^}]*\}", css)
    assert body and "white-space: pre-wrap" in body.group(0)


# -- the design pass (spec 5.3) ---------------------------------------------------------------------------

DETAILS_AT_START = 1   # app.js's only <details> when the design pass began: the Settings run-history log


def app_js() -> str:
    return (STATIC / "app.js").read_text(encoding="utf-8")


def css_block(css: str, header: str) -> str:
    """The body of the CSS block that starts with `header`, braces matched."""
    start = css.index(header)
    depth = 0
    for i in range(css.index("{", start), len(css)):
        if css[i] == "{":
            depth += 1
        elif css[i] == "}":
            depth -= 1
            if depth == 0:
                return css[start:i + 1]
    raise AssertionError("unbalanced braces after " + header)


def test_scorecard_group_order():
    """Every check, scorecard items and process rules together, grouped Missed, Unsure, Not applicable, Met."""
    panel = js_function(app_js(), "scorecardPanel")
    positions = [panel.index('"' + label + '"') for label in ("Missed", "Unsure", "Not applicable", "Met")]
    assert positions == sorted(positions), "the result groups are out of order"
    assert '"group-head"' in panel and "appendCheck(" in panel
    assert "detail.checks.forEach" not in panel, "checks are appended by group, not in one flat list"


def test_timeline_chips_class():
    css = (STATIC / "app.css").read_text(encoding="utf-8")
    assert re.search(r"\.timeline-strip\s*\{", css) and re.search(r"\.timeline-chip\s*\{", css)
    narrow = css_block(css, "@media (max-width: 640px)")
    strip = re.search(r"\.timeline-strip\s*\{[^}]*\}", narrow)
    assert strip and "flex-wrap: wrap" in strip.group(0), "the strip wraps at 640px"
    timeline = js_function(app_js(), "timelinePanel")
    assert '"timeline-strip"' in timeline and "timeline-chip" in timeline
    assert '"tl' not in timeline, "the old vertical timeline rows are gone"
    assert not re.search(r"\.tl(?![\w-])", css) and ".timeline {" not in css, "the old timeline styles are gone"


def test_no_new_collapsers():
    js = app_js()
    created = len(re.findall(r'h\(\s*"details"', js)) + len(re.findall(r'createElement\(\s*"details"', js))
    assert created <= DETAILS_AT_START


def test_tickets_layout():
    """One filter bar with Export, then the list beside the detail; the detail is three stacked panels."""
    js = app_js()
    render = js_function(js, "renderTickets")
    assert re.search(r'"panel-row cols-1-2"[^;]*tickets\.listBox, tickets\.detailBox', render), \
        "the list comes before the detail, so it sits above it when the row stacks"
    assert ".split2" not in (STATIC / "app.css").read_text(encoding="utf-8")
    assert 'var SINGLE_COLUMN = "(max-width: 1023px)"' in js, "matches the two-column stacking breakpoint"

    bar = js_function(js, "ticketFilterBar")
    for control in ('"tk-q"', '"tk-start"', '"tk-end"', '"tk-verdict"', '"tk-agent"', '"tk-missed"',
                    '"tk-priority"', '"Clear filters"', '"/api/export.xlsx"', '"Export to Excel"'):
        assert control in bar, control

    rows = js_function(js, "drawTicketList") + js_function(js, "ticketListRow")
    assert "/api/export.xlsx" not in rows, "Export moved to the filter bar"
    for kept in ("scoreText(row)", "row.agent_name", "dayLabel(row.date", "verdictBadge(row)", "row.missed",
                 "row.not_counted_reason", "recheckBadge()", '" sel"', '"data-ticket"', '"#" + row.ticket_id',
                 "capitalize(row.priority)", '"Show "', "TICKETS_PAGE", "location.hash = link"):
        assert kept in rows, kept

    detail = js_function(js, "drawDetail")
    order = [detail.index(part) for part in ('"dethead"', "notices,", "timelinePanel(", "scorecardPanel(",
                                              "commentsPanel(")]
    assert order == sorted(order), "the detail panels are out of spec order"
    assert "panel inner" not in detail
    for kept in ("Open in Zendesk", "Resolved on the call: ", "Review failed. ", "Not counted. ",
                 "The scores, outcomes and your overrides are kept.", "could not be confirmed"):
        assert kept in detail, kept

    comments = js_function(js, "commentsPanel")
    assert "panel inner" not in comments and "panels.push" not in comments, "one panel for every comment"
    assert '"Merged into ticket #"' in comments and '"group-head"' in comments


def run_result_groups(checks):
    """Run scorecardPanel's groupByResult on `checks` in node and return [label, count, [check ids]] rows."""
    node = shutil.which("node")
    if node is None:
        pytest.skip("node is not installed")
    script = (js_function(app_js(), "groupByResult") + "\nvar g = groupByResult(" + json.dumps(checks) + ");\n" +
              "process.stdout.write(JSON.stringify(g.map(function (x) { return [x.label, x.checks.length, " +
              "x.checks.map(function (c) { return c.check_id; })]; })));")
    result = subprocess.run([node, "-e", script], capture_output=True, text=True, check=True)
    return json.loads(result.stdout)


def test_result_groups_use_the_effective_outcome_in_rubric_order():
    def check(check_id, effective, outcome=None):
        return {"check_id": check_id, "effective": effective, "outcome": outcome or effective}

    checks = [check("sc01", "met"), check("sc02", "missed"), check("sc03", "met", "missed"),   # overridden
              check("sc04", "na"), check("sc05", "missed", "met"), check("pr01", "missed"), check("pr02", "met")]
    assert run_result_groups(checks) == [
        ["Missed", 3, ["sc02", "sc05", "pr01"]],
        ["Not applicable", 1, ["sc04"]],
        ["Met", 3, ["sc01", "sc03", "pr02"]],
    ]


def test_result_groups_keep_unfinished_checks_last():
    """A review error leaves checks with no outcome; they still show, in their own last group."""
    checks = [{"check_id": "sc01", "effective": None}, {"check_id": "pr01", "effective": "unsure"},
              {"check_id": "sc02", "effective": None}]
    assert run_result_groups(checks) == [["Unsure", 1, ["pr01"]], ["No result", 2, ["sc01", "sc02"]]]


# -- wide layout (phase 4, spec 4.3) ------------------------------------------------------------------

OWNER_WINDOW = 1920   # px, the owner's browser window
SCROLLBAR = 17        # px, the widest usual Windows scrollbar (Chrome's is 15)
DETAIL_BEFORE = 860   # px, about the Tickets detail's width at the owner's window before the wide layout
# px: the detail column the scorecard needs to fit without scrolling sideways. Its table's min-content width
# measured 533px on the demo data, plus the panel's padding (2 × 18px) and border (2 × 1px).
SCORECARD_DETAIL = 533 + 36 + 2


def cols_1_2_tracks(css: str, block: str) -> tuple[float, float, float, float]:
    """The (list min px, list fr, detail min px, detail fr) of `.panel-row.cols-1-2` in `block`."""
    from css_rules import px, values

    [columns] = values(block, ".panel-row.cols-1-2", "grid-template-columns")
    tracks = re.fullmatch(r"minmax\(([^,]+),\s*([\d.]+)fr\)\s+minmax\(([^,]+),\s*([\d.]+)fr\)", columns)
    assert tracks, columns
    list_min, list_fr, detail_min, detail_fr = tracks.groups()
    return px(css, list_min), float(list_fr), px(css, detail_min), float(detail_fr)


def detail_width(css: str, window: float, tracks: tuple[float, float, float, float]) -> float:
    """The detail column's width in a `window`-px browser window, from the page tokens and `tracks`."""
    from css_rules import px

    list_min, list_fr, _, detail_fr = tracks
    page = min(window - SCROLLBAR, px(css, "var(--page-max)"))
    content = page - 2 * px(css, "var(--pad-page-x)") - px(css, "var(--gap-panel)")
    return content - max(list_min, content * list_fr / (list_fr + detail_fr))


def test_ticket_list_column_wider():
    """Spec 4.3: from a breakpoint above 1024px the list gets at least 440px. The breakpoint sits where the
    detail still fits its scorecard beside a 440px list; below it (down to the 1024px stacking point) the
    list keeps today's 1:2 split. At the owner's window the detail is no narrower than it was."""
    from css_rules import app_css, media_block, top_level

    css = app_css()
    today = cols_1_2_tracks(css, top_level(css))
    assert today == (0, 1, 0, 2), "1024px up to the breakpoint: today's 1:2 split"
    assert detail_width(css, 1024, today) >= SCORECARD_DETAIL

    queries = [q for q in re.findall(r"@media (\(min-width: (\d+)px\))", css)
               if ".panel-row.cols-1-2" in media_block(css, q[0])]
    assert len(queries) == 1, "one wider-list block"
    query, breakpoint = queries[0][0], int(queries[0][1])
    wide = cols_1_2_tracks(css, media_block(css, query))
    assert wide[0] >= 440, "the list's minimum from the breakpoint up"
    assert breakpoint > 1024
    assert detail_width(css, breakpoint, wide) >= SCORECARD_DETAIL, "the scorecard fits at the breakpoint"
    assert detail_width(css, OWNER_WINDOW, wide) >= max(DETAIL_BEFORE, wide[2])


ROW_RUN_JS = """
globalThis.location = { hash: "" };
function shape(el) {
  if (!el.tagName) {
    return el.textContent;
  }
  return { tag: el.tagName, cls: el.className, kids: el.childNodes.map(shape) };
}
var out = ROWS.map(function (item) {
  var tr = ticketListRow(item.row, "#/tickets?ticket=" + item.row.ticket_id, item.selected);
  tr.listeners.click.forEach(function (f) { f({ target: { closest: function () { return null; } } }); });
  return { cls: tr.className, ticket: tr.getAttribute("data-ticket"), hash: location.hash,
           cells: tr.children.map(shape) };
});
process.stdout.write(JSON.stringify(out));
"""


def run_ticket_rows(rows):
    """Build Tickets list rows with app.js's ticketListRow in node. `rows` are {row, selected} items."""
    js = app_js()
    names = ("h", "append", "num", "localDate", "dayLabel", "capitalize", "shortError", "verdictBadge", "scoreText",
             "recheckBadge", "ticketListRow")
    source = "\n".join([js_var(js, v) for v in ("VERDICT_TEXT", "RECHECK_TITLE", "ERROR_TEXT_MAX")] +
                       [js_function(js, name) for name in names])
    return run_node(node_prelude() + source + f"\nvar ROWS = {json.dumps(rows)};\n" + ROW_RUN_JS)


def text_of(shape) -> str:
    return shape if isinstance(shape, str) else "".join(text_of(k) for k in shape["kids"])


def test_ticket_list_row_puts_the_verdict_on_the_first_line():
    """Spec 4.3: a compact row. Line one holds the score, ticket, agent, date, priority and verdict badge;
    the misses (or a not-counted call's reason) wrap under it."""
    base = {"ticket_id": 501, "agent_name": "Sam Okafor", "date": "2026-10-05", "priority": "normal",
            "counted": True, "review_error": False, "percent": 72.5, "verdict": "problem", "needs_recheck": True,
            "missed": ["10-minute rule", "#9 Resolution"], "not_counted_reason": None}
    missed, clean, left_out = run_ticket_rows([
        {"row": base, "selected": True},
        {"row": dict(base, ticket_id=502, priority=None, verdict="good", missed=[], needs_recheck=False,
                     percent=100), "selected": False},
        {"row": dict(base, ticket_id=503, counted=False, verdict="not_counted", percent=None,
                     not_counted_reason="Wrong number"), "selected": False},
    ])

    score, call = missed["cells"]
    assert "trow-score" in score["cls"].split() and text_of(score) == "72.5"
    first, second = call["kids"]
    assert first["cls"] == "trow-line"
    line = text_of(first)
    for part in ("#501", "Sam Okafor", "Oct 5", "Normal", "✗ Problem", "Recheck"):
        assert part in line, part
    assert line.index("#501") < line.index("Sam Okafor") < line.index("Oct 5") < line.index("Normal") \
        < line.index("✗ Problem")
    assert second["cls"] == "trow-missed" and text_of(second) == "Missed: 10-minute rule, #9 Resolution"
    assert missed["cls"].split() == ["rowlink", "sel"] and missed["ticket"] == "501"
    assert missed["hash"] == "#/tickets?ticket=501", "a click on the row opens the call"

    [only] = clean["cells"][1]["kids"]
    assert "✓ Good" in text_of(only) and "Normal" not in text_of(only), "no misses: one line"
    assert clean["cls"] == "rowlink"

    first, second = left_out["cells"][1]["kids"]
    assert "Not counted" in text_of(first) and "Recheck" not in text_of(first)
    assert text_of(second) == "Wrong number" and text_of(left_out["cells"][0]) == "-"


def test_review_failed_badge_and_filter_option():
    """T-045: a call whose review did not finish shows a Review failed badge in the failure color, with the
    error as its tooltip, and Review failed is a choice of the Verdict filter."""
    from css_rules import app_css, top_level, values

    row = {"ticket_id": 601, "agent_name": "Sam Okafor", "date": "2026-10-05", "priority": "normal",
           "counted": True, "review_error": True, "error": "request errored: overloaded_error", "percent": None,
           "verdict": "review_failed", "needs_recheck": False, "missed": [], "not_counted_reason": None}
    [built] = run_ticket_rows([{"row": row, "selected": False}])
    line = built["cells"][1]["kids"][0]
    badges = [k for k in line["kids"] if isinstance(k, dict) and k["cls"] == "trow-verdict"][0]["kids"]
    assert badges[0]["cls"] == "verdict review_failed" and text_of(badges[0]) == "! Review failed"
    assert text_of(built["cells"][0]) == "-"
    assert values(top_level(app_css()), ".verdict.review_failed", "color") == ["var(--failure)"]
    choices = js_var(app_js(), "VERDICT_CHOICES")
    assert '{ value: "review_failed", text: "Review failed" }' in choices


def test_ticket_list_row_styles():
    from css_rules import app_css, top_level, values

    css = top_level(app_css())
    assert values(css, ".tlist .trow-line", "flex-wrap") == ["wrap"], "line one wraps only when it must"
    assert values(css, ".tlist .trow-missed", "overflow-wrap") == ["anywhere"], "long misses wrap in the row"
    assert values(css, ".atable tr.sel td", "background") == ["var(--sel-tint)"], "the selection highlight stays"


# -- the reviewing model on each call (Task 4, T-034) ---------------------------------------------------

DETAIL_RUN_JS = """
globalThis.location = { hash: "" };
function timelinePanel() { return null; }
function scorecardPanel() { return null; }
function commentsPanel() { return null; }
var out = DETAILS.map(function (detail) {
  var tickets = { detailBox: document.createElement("div") };
  globalThis.tickets = tickets;
  drawDetail(detail);
  var head = tickets.detailBox.findAll(function (e) { return e.className === "dethead"; })[0];
  return head.findAll(function (e) { return e.className === "note"; }).map(function (e) { return e.textContent; });
});
process.stdout.write(JSON.stringify(out));
"""


def run_detail_header(details):
    """The call detail header's note lines, built by app.js's drawDetail in node for each API detail."""
    js = app_js()
    names = ("h", "append", "panel", "num", "pct", "localDate", "dayLabel", "capitalize", "shortError",
             "verdictBadge", "priorityText", "modelEffortText", "reviewedByText", "reviewAgainButton", "drawDetail")
    source = "\n".join([js_var(js, v) for v in ("VERDICT_TEXT", "RESOLVED_TEXT", "ERROR_TEXT_MAX")] +
                       [js_function(js, name) for name in names])
    return run_node(node_prelude() + source + f"\nvar DETAILS = {json.dumps(details)};\n" + DETAIL_RUN_JS)


def test_detail_shows_model(conn, version, tmp_path):
    sonnet = add_call(conn, version, make_call(130), all_met(version))
    conn.execute("UPDATE reviews SET model = 'claude-sonnet-5-5', effort = 'high' WHERE review_id = ?", (sonnet,))
    add_call(conn, version, make_call(131, start="2026-10-06T15:00:00Z"), all_met(version))
    conn.commit()
    conn.close()
    api = client(tmp_path)    # opening the database migrates the review without a model: Opus, effort NULL
    first, migrated = api.get("/api/tickets/130").json(), api.get("/api/tickets/131").json()
    assert (first["review"]["model_label"], first["review"]["effort"]) == ("Claude Sonnet 5.5", "high")
    assert (migrated["review"]["model_label"], migrated["review"]["effort"]) == ("Claude Opus 5.5", None)
    shown = run_detail_header([first, migrated])
    assert shown[0][-1] == "Reviewed by Claude Sonnet 5.5 · effort high"
    assert shown[1][-1] == "Reviewed by Claude Opus 5.5 · effort not recorded"


def test_timeline_and_merged_comments_use_minutes_and_the_collection_window(conn, version, tmp_path):
    """T-039: the timeline names the deadline in minutes (hours when whole), and the merged-into ticket's
    comments are described by the fixed collection window, not the deadline."""
    merged = MergeInfo(ticket_id=900, subject="Same site, earlier ticket", location="DEMO-104",
                       created_at="2026-10-04T12:00:00Z",
                       comments=(comment(9, "2026-10-05T17:00:00Z", "summary_email", "Merged reply", public=True),))
    add_call(conn, version, make_call(103, merged=merged), all_met(version))
    detail = client(tmp_path).get("/api/tickets/103").json()
    js = app_js()
    names = ("h", "append", "panel", "plural", "clockLabel", "minutesBetween", "formatMinutes", "timelinePanel",
             "commentsPanel")
    # The comment list itself isn't under test here.
    source = "\n".join(js_function(js, name) for name in names) + "\nfunction commentList() { return []; }\n"
    out = run_node(node_prelude() + source + f"\nvar DETAIL = {json.dumps(detail)};\n" + """
function texts(el) { return el.findAll(function (e) { return e.tagName === "LI"; }).map(function (e) { return e.textContent; }); }
var hours = JSON.parse(JSON.stringify(DETAIL));
hours.timeline.email_deadline_minutes = 1440;
var sub = commentsPanel(DETAIL, "").findAll(function (e) { return e.classList.contains("sub"); });
process.stdout.write(JSON.stringify({
  minutes: texts(timelinePanel(DETAIL)), hours: texts(timelinePanel(hours)),
  subs: sub.map(function (e) { return e.textContent; }),
  format: [30, 60, 90, 1440].map(formatMinutes)
}));
""")
    assert any(t.endswith("Summary email deadline (30 min after the call)") for t in out["minutes"]), out["minutes"]
    assert any(t.endswith("Summary email deadline (24 h after the call)") for t in out["hours"]), out["hours"]
    assert any("Its comments from this call's start until 24 hours after the call ends, or until the call was "
               "reviewed if that came first." in s for s in out["subs"])
    assert not any("deadline" in s for s in out["subs"])
    assert out["format"] == ["30 min", "1 h", "1 h 30 min", "24 h"]
