"""The Tickets list as an Excel workbook (.xlsx).

The app writes into an in-memory buffer and sends it as a download, so nothing lands on disk on the
server side. Every text value is stored as text: agent names and priorities come from Zendesk, and a
value starting with "=" must never become a formula.
"""
from __future__ import annotations

from collections.abc import Iterable
from datetime import date
from pathlib import Path
from typing import IO, Any

from openpyxl import Workbook
from openpyxl.styles import Font

from work_assistant.web.queries import zendesk_url

COLUMNS = ("Ticket", "Date", "Tech", "Priority", "Score", "Percent", "Verdict", "Missed", "Zendesk link")
VERDICT_LABELS = {"good": "Good", "needs_review": "Needs review", "problem": "Problem",
                  "review_failed": "Review failed", "not_counted": "Not counted"}
WIDTHS = {"A": 10, "B": 12, "C": 24, "D": 10, "E": 8, "F": 9, "G": 14, "H": 40, "I": 46}
SHEET_TITLE = "Tickets"


def verdict_label(row: dict[str, Any]) -> str:
    """What the Tickets list shows: a review error is "Review failed", never a score-based verdict."""
    if row.get("counted", True) and row.get("review_error"):
        return VERDICT_LABELS["review_failed"]
    return VERDICT_LABELS.get(row["verdict"], row["verdict"])


def _text(cell: Any, value: str | None) -> None:
    cell.value = value
    if value is not None:
        cell.data_type = "s"   # openpyxl would otherwise store "=..." as a formula


def write_xlsx(rows: Iterable[dict[str, Any]], path_or_buffer: str | Path | IO[bytes], *,
               subdomain: str | None = None) -> None:
    """Write `rows` (as `queries.contact_rows` returns them) to `path_or_buffer`, one row per call
    under the header `COLUMNS`. Score and Percent are blank for a review error or a Not counted call;
    Missed lists the short check labels; the Zendesk link is blank until the subdomain is set."""
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_TITLE
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, row in enumerate(rows, start=2):
        scored = row.get("counted", True) and not row.get("review_error")
        sheet.cell(index, 1, int(row["ticket_id"]))
        day = sheet.cell(index, 2, date.fromisoformat(row["date"]))
        day.number_format = "yyyy-mm-dd"
        _text(sheet.cell(index, 3), row.get("agent_name"))
        _text(sheet.cell(index, 4), row["priority"].capitalize() if row.get("priority") else None)
        sheet.cell(index, 5, row.get("earned") if scored else None)
        sheet.cell(index, 6, row.get("percent") if scored else None)
        _text(sheet.cell(index, 7), verdict_label(row))
        _text(sheet.cell(index, 8), ", ".join(row.get("missed") or []) or None)
        link = zendesk_url(subdomain, row["ticket_id"])
        cell = sheet.cell(index, 9)
        _text(cell, link)
        if link:
            cell.hyperlink = link
            cell.style = "Hyperlink"
    for column, width in WIDTHS.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    book.save(path_or_buffer)
