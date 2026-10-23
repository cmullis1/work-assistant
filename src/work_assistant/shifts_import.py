"""Import from Teams Shifts (Shifts import design, 2026-10-21-shift-finder-design.md). The department schedules
everyone in Microsoft Teams Shifts, whose "Export schedule" writes a date range to an Excel workbook. This module reads
those workbooks (`read_workbook`: bytes in, rows out) and plans what an import writes for each tech on the roster
(`plan_import`: it reads the database and writes nothing): their dated regions, their shift week by week, their
Started and their leave. `preview_json` is the plan as the Settings panel shows it, and `apply_import` writes it (Apply
backs up the database first, `backups.copy_beside`). The files are only ever read in memory, never kept.

The export's times are on the app's own zone (the exporter's clock). A Shifts row's day is the date its start falls on,
on the tech's region clock that day (its Group's region): an APAC shift starting Chicago Oct 7 15:00 starts Manila
Oct 8 04:00, so it is that tech's Oct 8. All-day Time Off entries (00:00 to 00:00) keep their written dates; a timed
one covers each day, on the tech's clock, that it overlaps by six hours or more (Ruling 16: a 9-hour entry is its one
day, one from Monday 09:00 to Wednesday 18:00 is Monday to Wednesday), and one under six hours in all is a half day, a
work day.

Leave is read as the owner's exports write it ("Settled with the owner after reading the real exports"): by word, from
a Shifts row's label and its notes, or a Time Off row's reason (`_kind`). A half day is a work day; a label that starts
with a time keeps only its own words from making it leave, not its notes' (Ruling 17).
"""
from __future__ import annotations

import io
import re
import sqlite3
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, replace
from datetime import date, datetime, time, timedelta, timezone
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from work_assistant import regions, settings
from work_assistant.activity import report, store, sync
from work_assistant.activity.model import ShiftVersion

SHIFTS_SHEET, TIME_OFF_SHEET = "Shifts", "Time Off"
NOT_AN_EXPORT = "{name} isn't a Teams Shifts export (no Shifts sheet)."
NO_SHIFTS = "The files have no shifts to import."
NOTE = "Teams Shifts"   # the note on the Time off entries an import writes: a later import replaces them
GROUPS = {"nasa team": "NASA", "emea team": "EMEA", "apac team": "APAC"}   # a row's Group, trimmed and lowercased
REST_DAY = "rest day"   # a label or reason, trimmed and lowercased, that makes the day a day off
# The words that make a day leave, case-insensitive, in a label, its notes or a reason: a sick word first, so a day
# with both ("Sick Leave") is Sick.
LEAVE_WORDS = (("sick", re.compile(r"\bsick|\bsl\b", re.IGNORECASE)),
               ("pto", re.compile(r"\bleave|\bpto\b|\bholiday|\bin\s+lieu|\bvacation|\bunpaid", re.IGNORECASE)))
# "half" at a word's start ("Half day", "Halfday", "Half-day"), not inside one ("on behalf of"), in a label or its
# notes: a half day, which is a work day.
HALF_DAY = re.compile(r"\bhalf", re.IGNORECASE)
# "9AM - 6PM DXB", "730AM - 230PM (Ramadan)", "14:00 - 18:00"
STARTS_WITH_TIME = re.compile(r"\d{1,2}(:?\d{2})?\s*[ap]m\b|\d{1,2}:\d{2}\b", re.IGNORECASE)
# A timed Time Off entry covers each day it overlaps this long or longer; one shorter in all is a half day.
FULL_DAY_OFF = timedelta(hours=6)
WEEKDAYS = ("Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat")
DAY_MINUTES = 24 * 60


@dataclass(frozen=True)
class ShiftRow:
    """One row of the Shifts sheet: a tech's shift on a day, or their rest day or leave, as its `label` (Custom
    Label) and `note` (Notes) say; each trimmed, "" when blank. `start` and `end` are aware, on the app's zone."""

    email: str
    name: str
    group: str
    start: datetime
    end: datetime
    label: str
    note: str


@dataclass(frozen=True)
class OffRow:
    """One row of the Time Off sheet that is a rest day or leave: `kind` is "pto", "sick" or "rest". An all-day entry
    covers its written dates, `first` through `last`. A timed one (six hours or more in all) covers each day it
    overlaps by six hours or more on the tech's clock (`_off_days`): it keeps its `start` and `end` (aware, on the
    app's zone), and `first` and `last` are the days they fall on, on the app's zone (an end at midnight, the day
    before)."""

    email: str
    name: str
    first: date
    last: date
    kind: str
    start: datetime | None = None
    end: datetime | None = None


@dataclass(frozen=True)
class TechPlan:
    """What an import writes for one tech. `regions` is their whole region history as it will be (changes before and
    after the imported dates kept, those inside them from the files). `entries` are the shift entries the import
    writes, oldest first; with `replace_before`, the entries before the imported dates go and the first whole imported
    week's pattern is dated where they began (or at the first imported date). `started`, when set, is their first
    tracked day (a late joiner's first working or leave day), and the first entry is dated it. Only an import that
    reaches back to the stored Team activity's first day sets either (Ruling 18, `plan_import`). `pto` and `sick` are
    their leave days. `problems` are what the preview flags for them."""

    email: str
    name: str
    regions: tuple[tuple[date, str], ...]
    entries: tuple[ShiftVersion, ...]
    replace_before: bool
    started: date | None
    pto: tuple[date, ...]
    sick: tuple[date, ...]
    problems: tuple[str, ...]


@dataclass(frozen=True)
class ImportPlan:
    """An import of some export files: the dates they cover (`first` to `last`, every Shifts row's day), each roster
    tech in them by region then name, the people skipped (not on the roster, by name), problems with no tech's plan
    to hold them, and `unrecognised`: the labels the import doesn't know on the rows it reads as working shifts
    (`_known`), each with how many rows have it, most common first, so the owner can check them."""

    first: date
    last: date
    techs: tuple[TechPlan, ...]
    skipped: tuple[str, ...]
    problems: tuple[str, ...]
    unrecognised: tuple[tuple[str, int], ...]


# -- the reader ---------------------------------------------------------------------------------------------------

SHIFT_COLUMNS = ("Member", "Work Email", "Group", "Start Date", "Start Time", "End Date", "End Time", "Custom Label",
                 "Notes")
TIME_OFF_COLUMNS = ("Member", "Work Email", "Start Date", "Start Time", "End Date", "End Time", "Time Off Reason")


def read_workbook(data: bytes, name: str, zone: str) -> tuple[list[ShiftRow], list[OffRow]]:
    """The Shifts rows and the rest-day and leave Time Off rows of an export workbook, with times on `zone` (the
    app's). Open shifts (no Work Email) and Time Off rows for anything else (Work From Home) are left out.
    ValueError (NOT_AN_EXPORT, with `name`) for a file that isn't an export workbook with a Shifts sheet."""
    refused = NOT_AN_EXPORT.format(name=name)
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError):
        raise ValueError(refused) from None
    try:
        if SHIFTS_SHEET not in book.sheetnames:
            raise ValueError(refused)
        tz = ZoneInfo(zone)
        shifts = [_shift_row(row, tz) for row in _records(book[SHIFTS_SHEET], SHIFT_COLUMNS)]
        offs = []
        if TIME_OFF_SHEET in book.sheetnames:
            offs = [off for row in _records(book[TIME_OFF_SHEET], TIME_OFF_COLUMNS)
                    if (off := _off_row(row, tz)) is not None]
    except (KeyError, TypeError, ValueError):
        raise ValueError(refused) from None
    finally:
        book.close()
    return shifts, offs


def _records(sheet, columns: tuple[str, ...]) -> list[dict]:
    """The sheet's rows that name a tech (a Work Email), each as {column: value} for `columns`, found by the header
    row's names. KeyError for a missing column."""
    sheet.reset_dimensions()   # read every row, whatever size the file says the sheet is
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or ()
    where = {str(title).strip(): i for i, title in enumerate(header) if title is not None}
    index = {column: where[column] for column in columns}
    found = []
    for row in rows:
        record = {column: row[i] if i < len(row) else None for column, i in index.items()}
        if _text(record["Work Email"]):
            found.append(record)
    return found


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _day(value) -> date:
    """A date cell: the export's M/D/YYYY text, or a real Excel date."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(_text(value), "%m/%d/%Y").date()


def _clock(value) -> time:
    """A time cell: the export's HH:MM text, or a real Excel time."""
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    return datetime.strptime(_text(value), "%H:%M").time()


def _at(record: dict, prefix: str, tz: ZoneInfo) -> datetime:
    return datetime.combine(_day(record[f"{prefix} Date"]), _clock(record[f"{prefix} Time"]), tzinfo=tz)


def _kind(label: str, note: str = "") -> str | None:
    """What a Shifts row's label and notes, or a Time Off row's reason, make the day: "rest", "pto" or "sick", or None
    for a working shift. `Rest Day` is a day off. A half day ("half" at a word's start, in the label or notes) is a
    working shift whatever else they say. Otherwise a sick word, then a leave word, in the label or the notes makes it
    leave; a label that starts with a time ("11AM - 6PM (Ramadan)") is only its hours, so only the notes' words count
    then (Ruling 17). Anything else ("Training", a client's name) is a working shift."""
    if label.lower() == REST_DAY:
        return "rest"
    if HALF_DAY.search(label) or HALF_DAY.search(note):
        return None
    texts = (note,) if STARTS_WITH_TIME.match(label) else (label, note)
    for kind, words in LEAVE_WORDS:
        if any(words.search(text) for text in texts):
            return kind
    return None


def _known(label: str) -> bool:
    """Whether the import knows a label by itself, without its notes: blank, a rest day, a half day, one that starts
    with a time, or one with a leave word. The labels it doesn't know are read as working shifts and listed."""
    return (not label or bool(HALF_DAY.search(label)) or bool(STARTS_WITH_TIME.match(label))
            or _kind(label) is not None)


def _shift_row(record: dict, tz: ZoneInfo) -> ShiftRow:
    email = _text(record["Work Email"]).lower()
    return ShiftRow(email=email, name=_text(record["Member"]) or email, group=_text(record["Group"]),
                    start=_at(record, "Start", tz), end=_at(record, "End", tz), label=_text(record["Custom Label"]),
                    note=_text(record["Notes"]))


def _off_row(record: dict, tz: ZoneInfo) -> OffRow | None:
    """A Time Off row as an OffRow, or None when its reason isn't a rest day or leave, or when it is a timed entry
    shorter than FULL_DAY_OFF in all (a half day, a work day). An all-day entry covers its start date up to the day
    before its end date; a timed one keeps its start and end, and its days are found on the tech's clock
    (`_off_days`)."""
    kind = _kind(_text(record["Time Off Reason"]))
    if kind is None:
        return None
    email = _text(record["Work Email"]).lower()
    name = _text(record["Member"]) or email
    start, end = _at(record, "Start", tz), _at(record, "End", tz)
    last = max(start.date(), end.date() - timedelta(days=1) if end.time() == time(0) else end.date())
    if start.time() == time(0) and end.time() == time(0):
        return OffRow(email, name, start.date(), last, kind)
    if end - start < FULL_DAY_OFF:
        return None
    return OffRow(email, name, start.date(), last, kind, start, end)


# -- the planner --------------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class _Placed:
    """A Shifts row on its tech's clock: the clock of its Group's region (the app's zone when the Group isn't one)."""

    row: ShiftRow
    region: str | None   # None: the row's Group isn't NASA Team, EMEA Team or APAC Team
    day: date
    start_min: int
    end_min: int         # minutes after the day's midnight: past DAY_MINUTES when it ends the next day
    kind: str | None     # "rest", "pto" or "sick"; None for a working shift


def _placed(row: ShiftRow, app_zone: str) -> _Placed:
    region = GROUPS.get(row.group.strip().lower())
    zone = ZoneInfo(regions.REGIONS[region] if region else app_zone)
    start, end = row.start.astimezone(zone), row.end.astimezone(zone)
    day = start.date()
    return _Placed(row, region, day, start.hour * 60 + start.minute,
                   (end.date() - day).days * DAY_MINUTES + end.hour * 60 + end.minute, _kind(row.label, row.note))


def _sunday(day: date) -> date:
    return day - timedelta(days=_weekday(day))


def _weekday(day: date) -> int:
    """Sunday = 0 .. Saturday = 6, as shift entries count them."""
    return (day.weekday() + 1) % 7


def plan_import(conn: sqlite3.Connection, shifts: list[ShiftRow], offs: list[OffRow]) -> ImportPlan:
    """What importing these rows would write for each roster tech in them (`settings.team_emails`); everyone else is
    skipped. The labels it doesn't know are counted on the rows it reads as their working shifts (in NASA Team, EMEA
    Team or APAC Team). Reads the database (each tech's regions, shift entries and name) and writes nothing.
    For each tech the import reaches back (Ruling 18) when the files' first date is on or before the first day the
    stored Team activity covers on that tech's clock (`sync.first_covered_day`), or when no activity is stored yet:
    only then may it replace the entries before the imported dates and set a Started, so next month's export leaves
    the history before it as it is. ValueError (NO_SHIFTS) when there are no Shifts rows."""
    if not shifts:
        raise ValueError(NO_SHIFTS)
    app_zone = settings.get_timezone(conn)
    placed = sorted((_placed(row, app_zone) for row in shifts), key=lambda p: (p.day, p.row.start))
    first, last = placed[0].day, placed[-1].day
    first_week = min((_sunday(p.day) for p in placed if p.kind is None), default=None)
    covered = store.synced_from(conn)
    roster = set(settings.team_emails(conn))
    known = {tech["email"]: tech["name"] for tech in report.techs(conn)}
    skipped = sorted({row.name for row in [*shifts, *offs] if row.email not in roster}, key=lambda n: (n.lower(), n))
    techs, problems = [], []
    for email in sorted({row.email for row in [*shifts, *offs] if row.email in roster}):
        own = [p for p in placed if p.row.email == email]
        name = known.get(email, email)
        if name == email:   # the app has no name for them: the file's
            name = own[0].row.name if own else next(off.name for off in offs if off.email == email)
        if not any(p.region for p in own):
            problems.append(f"{name} has no shifts in NASA Team, EMEA Team or APAC Team, so nothing is imported "
                            "for them.")
            continue
        reaches_back = covered is None or first <= sync.first_covered_day(conn, covered, email)
        techs.append(_plan_tech(conn, email, name, own, [off for off in offs if off.email == email],
                                first, last, first_week, reaches_back))
    techs.sort(key=lambda t: (regions.ORDER.index(t.regions[-1][1]), t.name.lower(), t.email))
    unknown = Counter(p.row.label for p in placed if p.row.email in roster and p.region and p.kind is None
                      and not _known(p.row.label))
    unrecognised = sorted(unknown.items(), key=lambda item: (-item[1], item[0].lower(), item[0]))
    return ImportPlan(first, last, tuple(techs), tuple(skipped), tuple(problems), tuple(unrecognised))


def _plan_tech(conn: sqlite3.Connection, email: str, name: str, own: list[_Placed], offs: list[OffRow],
               first: date, last: date, first_week: date | None, reaches_back: bool) -> TechPlan:
    """One tech's plan from their rows (`own`, by day), over the imported dates `first` to `last`. `first_week` is
    the first week anyone has a working shift in the files. Only an import that `reaches_back` (Ruling 18) replaces
    the entries before the imported dates (when they hold one pattern or none, else it keeps them and says so) and
    sets a late joiner's Started; one that doesn't keeps those entries and the tech's Started exactly as they are."""
    problems = [f'Skipped {count} of {name}\'s rows in group "{group}": only NASA Team, EMEA Team and APAC Team are '
                "imported." for group, count in Counter(p.row.group for p in own if p.region is None).items()]
    rows = [p for p in own if p.region]
    history = _regions(regions.region_history(conn, email), rows, first, last)
    working = [p for p in rows if p.kind is None]
    pto, sick = _leave(rows, offs, history)
    # Leave counts as a scheduled work day of its week (Ruling 12); a Rest Day stays a day off when leave spans it.
    rest = {p.day for p in rows if p.kind == "rest"} | {day for off in offs if off.kind == "rest"
                                                         for day in _off_days(off, history)}
    leave = sorted(day for day in {*pto, *sick} - rest if first <= day <= last)
    entries = _weekly(email, working, leave, first, problems)
    before = [v for v in store.shift_versions(conn, email) if v.effective_from < first]
    replace_before = reaches_back and bool(entries) and len({_pattern(v) for v in before}) <= 1
    started = None
    if replace_before and first_week is not None:
        scheduled = min([working[0].day, *leave[:1]])   # their first working or leave day
        if scheduled > first_week + timedelta(days=6):
            started = scheduled
    elif reaches_back and entries:   # the history before is kept, and a Started would delete it
        problems.append(f"Kept the shift history before {first.isoformat()}: "
                        + "; ".join(f"from {v.effective_from.isoformat()} {_described(v)}" for v in before) + ".")
    if replace_before:
        _date_back(entries, started or (before[0].effective_from if before else first))
    return TechPlan(email, name, history, tuple(entries), replace_before, started, pto, sick, tuple(problems))


def _date_back(entries: list[ShiftVersion], since: date) -> None:
    """Date the tech's first entry `since`: where the entries before the imported dates began, or their Started. The
    pattern dated back is the first whole imported week's, the first week whose Sunday is on or after its first date:
    a first entry dated mid-week, where the first imported date cuts its week, holds only that week's days from then
    on (Wed-Fri for a Mon-Fri tech when the files begin on a Wednesday). That cut entry then keeps its own dates when
    its pattern differs. With no whole week in the files, the cut week's is used."""
    head = entries[0]
    whole = head.effective_from + timedelta(days=(7 - _weekday(head.effective_from)) % 7)
    pattern = [v for v in entries if v.effective_from <= whole][-1]   # the entry in force on that week's Sunday
    if since < head.effective_from and pattern is not head:
        entries.insert(0, replace(pattern, effective_from=since))
    else:
        entries[0] = replace(head, effective_from=since)


def _regions(stored: list[tuple[date, str]], rows: list[_Placed], first: date,
             last: date) -> tuple[tuple[date, str], ...]:
    """The tech's region history as the import leaves it: `stored` changes before `first` and after `last` kept,
    and inside them one change per run of the same region in `rows`, dated at the run's first day."""
    runs: list[tuple[date, str]] = []
    for p in rows:
        if runs and runs[-1][1] == p.region:
            continue
        if runs and runs[-1][0] == p.day:
            runs[-1] = (p.day, p.region)   # two regions on one day: the day's later row decides it
        else:
            runs.append((p.day, p.region))
    kept = [(since, region) for since, region in stored if since < first or since > last]
    return tuple(regions.checked(kept + runs))


def _weekly(email: str, working: list[_Placed], leave: list[date], first: date,
            problems: list[str]) -> list[ShiftVersion]:
    """The tech's shift entries from their working shifts (by day) and their leave days (`leave`, sorted), one per
    week from their first week with either to their last, repeats merged. A leave day is a work day of its week
    (Ruling 12), so a week with leave is no shift change. A week with working shifts: their weekdays and its leave
    days, at the shifts' most common hours (the earliest day's on a tie). A week with leave but no working shift:
    its leave days, at the week before's hours (with no week before, the next working week's). Each dated its
    Sunday, or `first` when that is later. A week with neither: no work days, the week before's hours. A week with
    a shift that crosses midnight on the tech's clock is left as the week before, and `problems` says so."""
    if not working:
        return []
    weeks: dict[date, list[_Placed]] = {}
    crossing = set()
    for p in working:
        weeks.setdefault(_sunday(p.day), []).append(p)
        if p.end_min > DAY_MINUTES:
            crossing.add(_sunday(p.day))
            problems.append(f"The {report.clock(p.start_min)}-{report.clock(p.end_min - DAY_MINUTES)} shift on "
                            f"{p.day.isoformat()} crosses midnight on {regions.CITIES[p.region]} time, so its week "
                            "keeps the week before's shift.")
    off: dict[date, set[int]] = {}
    for day in leave:
        off.setdefault(_sunday(day), set()).add(_weekday(day))

    def hours(week: date) -> tuple[int, int]:
        return Counter((p.start_min, p.end_min) for p in weeks[week]).most_common(1)[0][0]

    working_weeks = sorted(week for week in weeks if week not in crossing)
    entries: list[ShiftVersion] = []
    week, final = min([*weeks, *off]), max([*weeks, *off])
    while week <= final:
        days_off = frozenset(off.get(week, ()))
        if week in crossing:
            pass
        elif week in weeks:
            days = frozenset(_weekday(p.day) for p in weeks[week]) | days_off
            _add(entries, ShiftVersion(email, max(week, first), days, *hours(week), True))
        elif days_off and entries:
            _add(entries, replace(entries[-1], effective_from=week, work_days=days_off))
        elif days_off:
            later = next((w for w in working_weeks if w > week), None)
            if later is not None:
                _add(entries, ShiftVersion(email, max(week, first), days_off, *hours(later), True))
        elif entries:
            _add(entries, replace(entries[-1], effective_from=week, work_days=frozenset()))
        week += timedelta(days=7)
    return entries


def _pattern(version: ShiftVersion) -> tuple:
    return version.work_days, version.start_min, version.end_min, version.active


def _add(entries: list[ShiftVersion], version: ShiftVersion) -> None:
    """Add `version` unless it repeats the entry before it."""
    if not entries or _pattern(entries[-1]) != _pattern(version):
        entries.append(version)


def _leave(rows: list[_Placed], offs: list[OffRow],
           history: tuple[tuple[date, str], ...]) -> tuple[tuple[date, ...], tuple[date, ...]]:
    """The tech's PTO days and Sick days, from their Shifts rows that are leave and their leave Time Off rows, each a
    whole day; a day that is both PTO and Sick is Sick."""
    found = {"pto": set(), "sick": set()}
    for p in rows:
        if p.kind in found:
            found[p.kind].add(p.day)
    for off in offs:
        if off.kind in found:
            found[off.kind].update(_off_days(off, history))
    pto = found["pto"] - found["sick"]
    return tuple(sorted(pto)), tuple(sorted(found["sick"]))


def _off_days(off: OffRow, history: tuple[tuple[date, str], ...]) -> list[date]:
    """The days a Time Off row covers: an all-day one's written dates; a timed one's, each day on the tech's clock
    (their region's at each moment of it, `_region_at`) that it overlaps by FULL_DAY_OFF or more (Ruling 16). A
    9-hour entry is its one day: an APAC one from Chicago 15:00 to midnight is Manila's next day, 04:00 to 13:00."""
    if off.start is None:
        return _dates(off.first, off.last)
    begins = [datetime.combine(since, time(0), tzinfo=ZoneInfo(regions.REGIONS[region])).astimezone(timezone.utc)
              for since, region in history]
    moment, end = off.start.astimezone(timezone.utc), off.end.astimezone(timezone.utc)
    spent: dict[date, timedelta] = defaultdict(timedelta)
    while moment < end:   # one stretch at a time on one day of one region's clock
        zone = ZoneInfo(regions.REGIONS[_region_at(history, moment)])
        day = moment.astimezone(zone).date()
        midnight = datetime.combine(day + timedelta(days=1), time(0), tzinfo=zone).astimezone(timezone.utc)
        until = min([end, midnight, *(begun for begun in begins if begun > moment)])
        spent[day] += until - moment
        moment = until
    return sorted(day for day, length in spent.items() if length >= FULL_DAY_OFF)


def _region_at(history: tuple[tuple[date, str], ...], moment: datetime) -> str:
    """The tech's region at `moment`: that of the latest change in `history` whose day has begun on its own region's
    clock (the earliest change's before any). A change to APAC dated Oct 8 begins at Manila's midnight, Chicago's
    Oct 7 11:00, so Chicago's Oct 7 15:00 is on Manila's clock, and is that tech's Oct 8."""
    region = history[0][1]
    for since, change in history:
        if datetime.combine(since, time(0), tzinfo=ZoneInfo(regions.REGIONS[change])) <= moment:
            region = change
    return region


def _dates(first: date, last: date) -> list[date]:
    return [first + timedelta(days=i) for i in range((last - first).days + 1)]


def _described(version: ShiftVersion) -> str:
    """A shift entry as the preview writes it: "Mon-Fri 08:00-17:00" (a run of three or more days) or
    "Sun,Tue 08:00-17:00"."""
    days = sorted(version.work_days)
    if len(days) > 2 and days == list(range(days[0], days[-1] + 1)):
        text = f"{WEEKDAYS[days[0]]}-{WEEKDAYS[days[-1]]}"
    else:
        text = ",".join(WEEKDAYS[d] for d in days) or "no days"
    return f"{text} {report.clock(version.start_min)}-{report.clock(version.end_min)}"


# -- the preview --------------------------------------------------------------------------------------------------


def preview_json(plan: ImportPlan) -> dict:
    """The plan as the Settings panel shows it, JSON-ready: dates ISO, hours "HH:MM", days 0 (Sunday) to 6."""
    return {
        "first": plan.first.isoformat(),
        "last": plan.last.isoformat(),
        "techs": [{
            "email": tech.email,
            "name": tech.name,
            "regions": [{"since": since.isoformat(), "region": region} for since, region in tech.regions],
            "changes": [{"since": v.effective_from.isoformat(), "days": sorted(v.work_days),
                         "start": report.clock(v.start_min), "end": report.clock(v.end_min)} for v in tech.entries],
            "started": tech.started.isoformat() if tech.started else None,
            "replace_before": tech.replace_before,
            "pto": [day.isoformat() for day in tech.pto],
            "sick": [day.isoformat() for day in tech.sick],
            "problems": list(tech.problems),
        } for tech in plan.techs],
        "skipped": list(plan.skipped),
        "problems": list(plan.problems),
        "unrecognised": [{"label": label, "count": count} for label, count in plan.unrecognised],
    }


# -- the applier --------------------------------------------------------------------------------------------------


class _OneTransaction:
    """The connection as the store's writers see it while one tech's plan is written. Each writer commits its own
    change (`conn.commit()`, or a `with conn:` block); here those commits wait, so `apply_import` commits the tech's
    writes together, or rolls them all back."""

    def __init__(self, conn: sqlite3.Connection) -> None:
        self._conn = conn

    def __getattr__(self, name: str):
        return getattr(self._conn, name)

    def commit(self) -> None:
        """Held: the tech's writes are committed together."""

    def __enter__(self) -> _OneTransaction:
        return self

    def __exit__(self, *exc) -> None:
        return None


def apply_import(conn: sqlite3.Connection, plan: ImportPlan) -> dict:
    """Write `plan`, each tech in one transaction: their region history (`regions.set_regions`); their shift
    entries, in place of those dated inside the imported dates (and those before them, with `replace_before`), then
    their Started (`store.set_start`) and repeats merged; and their leave as Time off entries noted NOTE
    (`_write_time_off`). Returns how many techs it wrote, and how many shift entries and NOTE Time off entries the
    files give them."""
    written = {"techs": 0, "entries": 0, "time_off": 0}
    for tech in plan.techs:
        try:
            time_off = _write_tech(_OneTransaction(conn), plan, tech)
        except BaseException:
            conn.rollback()
            raise
        conn.commit()
        written["techs"] += 1
        written["entries"] += len(tech.entries)
        written["time_off"] += time_off
    return written


def _write_tech(conn: _OneTransaction, plan: ImportPlan, tech: TechPlan) -> int:
    """Write one tech's plan; returns how many NOTE Time off entries they have for it."""
    regions.set_regions(conn, tech.email, list(tech.regions))
    if tech.entries:
        store.remove_shift_versions(conn, tech.email, date.min if tech.replace_before else plan.first, plan.last)
        for version in tech.entries:
            store.add_shift_version(conn, version)
        if tech.started:
            store.set_start(conn, tech.email, tech.started)   # repeats merged too
        else:
            store.merge_repeats(conn, tech.email)
    return _write_time_off(conn, plan, tech)


def _write_time_off(conn: _OneTransaction, plan: ImportPlan, tech: TechPlan) -> int:
    """The tech's PTO and Sick days as Time off entries noted NOTE, consecutive days of one kind as one entry; returns
    how many there are. An earlier import's entries on the imported dates (noted NOTE) are replaced, keeping their
    days outside the dates. A day any other entry covers gets none, and every other entry, a lead's own above all, is
    left as it is. An entry already as the files give it is kept, so the same files again change nothing."""
    earlier, covered = [], set()
    for entry in store.time_off(conn, tech.email):
        if entry.note == NOTE and entry.first <= plan.last and entry.last >= plan.first:
            earlier.append(entry)
        else:
            covered.update(_dates(entry.first, entry.last))
    kinds = {day: entry.kind for entry in earlier for day in _dates(entry.first, entry.last)
             if not plan.first <= day <= plan.last}
    kinds.update({day: kind for kind, days in (("pto", tech.pto), ("sick", tech.sick)) for day in days
                  if day not in covered})
    wanted = _runs(kinds)
    kept = set()
    for entry in earlier:
        run = (entry.first, entry.last, entry.kind)
        if run in wanted and run not in kept:
            kept.add(run)
        else:
            store.remove_time_off(conn, entry.id)
    for first, last, kind in wanted:
        if (first, last, kind) not in kept:
            store.add_time_off(conn, tech.email, first, last, NOTE, kind)
    return len(wanted)


def _runs(kinds: dict[date, str]) -> list[tuple[date, date, str]]:
    """Each run of consecutive days of one kind in `kinds` (day: kind), as (first, last, kind), oldest first."""
    runs: list[tuple[date, date, str]] = []
    for day in sorted(kinds):
        if runs and runs[-1][2] == kinds[day] and runs[-1][1] + timedelta(days=1) == day:
            runs[-1] = (runs[-1][0], day, kinds[day])
        else:
            runs.append((day, day, kinds[day]))
    return runs
